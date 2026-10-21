from pathlib import Path

import docx

from jobhunter.ingestion.extract_text import extract_text
from jobhunter.ingestion.loader import combined_text, load_inbox


def test_extract_text_txt(tmp_path: Path):
    f = tmp_path / "note.txt"
    f.write_text("Bachelor of Science, Chemistry, 2018", encoding="utf-8")
    assert "Chemistry" in extract_text(f)


def test_extract_text_docx(tmp_path: Path):
    f = tmp_path / "cv.docx"
    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Quality Control Specialist, 5 years experience")
    doc.save(str(f))

    text = extract_text(f)
    assert "Jane Doe" in text
    assert "Quality Control" in text


def test_load_inbox_walks_subfolders_and_skips_unsupported(tmp_path: Path):
    (tmp_path / "cv").mkdir()
    (tmp_path / "diplomas").mkdir()
    (tmp_path / "cv" / "resume.txt").write_text("resume content", encoding="utf-8")
    (tmp_path / "diplomas" / "degree.txt").write_text("BSc Biology", encoding="utf-8")
    (tmp_path / "ignore.zip").write_bytes(b"not a real zip")

    docs = load_inbox(tmp_path)
    names = {d.path.name for d in docs}
    assert names == {"resume.txt", "degree.txt"}


def test_load_inbox_missing_dir_returns_empty(tmp_path: Path):
    assert load_inbox(tmp_path / "does_not_exist") == []


def test_combined_text_labels_each_source(tmp_path: Path):
    a = tmp_path / "a.txt"
    a.write_text("Alpha content", encoding="utf-8")
    docs = load_inbox(tmp_path)
    text = combined_text(docs)
    assert "SOURCE: a.txt" in text
    assert "Alpha content" in text
